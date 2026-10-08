import json, sqlite3, uuid, shutil, io, re, os, tempfile
from pathlib import Path
from datetime import datetime
from functools import wraps
from flask import Flask, request, jsonify, send_file, send_from_directory
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from grn_parser import parse_file as parse_grn
from po_parser import parse_file as parse_po

ROOT=Path(__file__).resolve().parent
DATA=ROOT/'data'; ORIGINALS=DATA/'originals'; GENERATED=DATA/'generated'; DB=DATA/'control_center.db'
# Cloud Run's filesystem is ephemeral. Keep uploads in /tmp while parsing, then
# persist the original in Firebase Storage when cloud mode is enabled.
TMP_DIR=Path(os.getenv('TMPDIR','/tmp')) / 'sayza-po-grn'
for p in (ORIGINALS,GENERATED,TMP_DIR): p.mkdir(parents=True,exist_ok=True)
app=Flask(__name__,static_folder='static',static_url_path='')

# ---------- Firebase Authentication + Cloud storage ----------
USE_FIRESTORE=os.getenv('USE_FIRESTORE','1' if os.getenv('K_SERVICE') else '0') == '1'
_firebase_auth_ready=False
_firestore=None
_storage_bucket=None
_firebase_init_error=''

def init_firebase_cloud():
    """Initialize one Firebase Admin app and all cloud clients.

    The previous V2 code swallowed initialization errors and left
    ``_firestore`` as None, which later caused ``NoneType.collection()``
    errors.  Initialization is now explicit and fails early with the real
    Firebase error when cloud mode is enabled.
    """
    global _firebase_auth_ready, _firestore, _storage_bucket, _firebase_init_error, firebase_auth
    if not USE_FIRESTORE and _firebase_auth_ready and _firestore is not None:
        return
    try:
        import firebase_admin
        from firebase_admin import auth as _firebase_auth_module, firestore, storage
        firebase_auth = _firebase_auth_module

        try:
            firebase_admin.get_app()
        except ValueError:
            firebase_admin.initialize_app(options={
                'projectId': os.getenv('FIREBASE_PROJECT_ID','sayza-po-grn'),
                'storageBucket': os.getenv('FIREBASE_STORAGE_BUCKET','sayza-po-grn.firebasestorage.app'),
            })

        _firebase_auth_ready=True
        _firestore=firestore.client()
        _storage_bucket=storage.bucket()
        _firebase_init_error=''
    except Exception as exc:
        _firebase_auth_ready=False
        _firestore=None
        _storage_bucket=None
        _firebase_init_error=f'{type(exc).__name__}: {exc}'
        if USE_FIRESTORE:
            raise RuntimeError(f'Firebase cloud initialization failed: {_firebase_init_error}') from exc

init_firebase_cloud()


def require_auth(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if os.getenv('DEV_AUTH_BYPASS','0') == '1' and not os.getenv('K_SERVICE'):
            return fn(*args, **kwargs)
        if not _firebase_auth_ready:
            return jsonify({'error':'Authentication service is not available'}),503
        header=request.headers.get('Authorization','')
        if not header.startswith('Bearer '):
            return jsonify({'error':'Authentication required'}),401
        try:
            decoded=firebase_auth.verify_id_token(header[7:].strip())
            request.firebase_user=decoded
        except Exception:
            return jsonify({'error':'Invalid or expired authentication token'}),401
        return fn(*args, **kwargs)
    return wrapper

PO_COLUMNS=["Sr. No.","Item Code","Item Desc","HSN Code","PO Number","PO Date","PO Release Date","Payment Terms","Expected Delivery Date","PO Expiry Date","Vendor Name","Qty","MRP (INR)","Unit Base Cost (INR)","Taxable Value (INR)","CGST Rate","CGST Amount","SGST/UGST Rate","SGST/UGST Amount","IGST Rate","IGST Amount","CESS Rate","CESS Amount","Additional CESS Rate","Additional CESS Amount","Total (INR)"]
GRN_COLUMNS=["Sr. No.","SKU Code","SKU Desc","PO Number","GRN Number","Invoice Number","Lot MRP (INR)","Exp Qty","Recv Qty","Unit Price (INR)","Taxable Value (INR)","CGST Rate","CGST Amount","SGST/UGST Rate","SGST/UGST Amount","IGST Rate","IGST Amount","CESS Rate","CESS Amount","Add. Cess Rate","Add. Cess Amount","Total (INR)"]

# ---------- Local DB (kept for local development and migration source) ----------
def conn():
    c=sqlite3.connect(DB); c.row_factory=sqlite3.Row; return c

def init_local():
    c=conn()
    c.execute('''CREATE TABLE IF NOT EXISTS documents(id TEXT PRIMARY KEY, doc_type TEXT, filename TEXT, stored_name TEXT, uploaded_at TEXT, po_number TEXT, grn_number TEXT, invoice_number TEXT, vendor_name TEXT, meta_json TEXT, item_count INTEGER, status TEXT, error TEXT)''')
    c.execute('''CREATE TABLE IF NOT EXISTS lines(id TEXT PRIMARY KEY, doc_id TEXT, doc_type TEXT, data_json TEXT, row_index INTEGER DEFAULT 0)''')
    try: c.execute('ALTER TABLE lines ADD COLUMN row_index INTEGER DEFAULT 0')
    except sqlite3.OperationalError: pass
    c.execute('''CREATE TABLE IF NOT EXISTS appointments(po_number TEXT PRIMARY KEY, appointment_date TEXT, updated_at TEXT)''')
    c.commit(); c.close()
init_local()

def clean_num(v):
    try: return float(v or 0)
    except: return 0.0

def normalize_key(v): return str(v or '').strip().upper()

def fs_docs(): return _firestore.collection('documents')
def fs_lines(): return _firestore.collection('document_lines')
def fs_appointments(): return _firestore.collection('appointments')

def ensure_cloud():
    if not USE_FIRESTORE: return
    if not _firestore or not _storage_bucket: raise RuntimeError('Firestore/Storage is not available')

def save_doc(doc_type, f):
    did=uuid.uuid4().hex
    safe=f'{did}_{Path(f.filename).name}'.replace('/','_').replace('\\','_')
    tmp=TMP_DIR/(did+'_tmp_'+Path(f.filename).name)
    f.save(tmp)
    try:
        if doc_type=='PO': rows=parse_po(str(tmp)); meta={}; vendor=''
        else:
            rows,meta,vendor,text=parse_grn(str(tmp)); meta=meta or {}
            if vendor: meta['vendorName']=vendor
        if not rows: raise ValueError('No item rows could be extracted')
        if doc_type=='PO':
            po=rows[0].get('PO Number',''); grn=''; inv=''; vendor=rows[0].get('Vendor Name','')
        else:
            po=rows[0].get('PO Number',''); grn=rows[0].get('GRN Number',''); inv=rows[0].get('Invoice Number',''); vendor=meta.get('vendorName','')
        now=datetime.now().isoformat(timespec='seconds')
        storage_path=f'originals/{did}/{safe}'
        if USE_FIRESTORE:
            ensure_cloud()
            blob=_storage_bucket.blob(storage_path)
            blob.upload_from_filename(str(tmp),content_type=f.mimetype or 'application/pdf')
            doc={'id':did,'doc_type':doc_type,'filename':f.filename,'stored_name':safe,'storage_path':storage_path,'uploaded_at':now,'po_number':po,'grn_number':grn,'invoice_number':inv,'vendor_name':vendor,'meta_json':meta,'item_count':len(rows),'status':'Success','error':''}
            batch=_firestore.batch(); batch.set(fs_docs().document(did),doc)
            for idx,row in enumerate(rows):
                lid=uuid.uuid4().hex
                batch.set(fs_lines().document(lid),{'id':lid,'doc_id':did,'doc_type':doc_type,'data':row,'row_index':idx})
                if (idx+1)%400==0: batch.commit(); batch=_firestore.batch()
            batch.commit()
        else:
            shutil.move(str(tmp), ORIGINALS/safe)
            c=conn(); c.execute('INSERT INTO documents VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',(did,doc_type,f.filename,safe,now,po,grn,inv,vendor,json.dumps(meta),len(rows),'Success',''))
            for idx,row in enumerate(rows): c.execute('INSERT INTO lines(id,doc_id,doc_type,data_json,row_index) VALUES(?,?,?,?,?)',(uuid.uuid4().hex,did,doc_type,json.dumps(row,ensure_ascii=False),idx))
            c.commit(); c.close()
        return {'id':did,'filename':f.filename,'type':doc_type,'status':'Success','items':len(rows),'po_number':po,'grn_number':grn}
    except Exception as e:
        try: tmp.unlink(missing_ok=True)
        except Exception: pass
        if USE_FIRESTORE:
            try:
                if _storage_bucket:
                    _storage_bucket.blob(f'originals/{did}/{safe}').delete()
            except Exception:
                pass
        else:
            c=conn(); now=datetime.now().isoformat(timespec='seconds')
            c.execute('INSERT INTO documents VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',(did,doc_type,f.filename,'',now,'','','','', '{}',0,'Failed',str(e))); c.commit(); c.close()
        return {'id':did,'filename':f.filename,'type':doc_type,'status':'Failed','error':str(e)}

def cloud_documents(doc_type):
    docs=[]
    for s in fs_docs().where('doc_type','==',doc_type.upper()).stream():
        d=s.to_dict(); d.setdefault('id',s.id); docs.append(d)
    docs.sort(key=lambda x:x.get('uploaded_at',''), reverse=True)
    if doc_type.upper()=='PO':
        appts={s.id:s.to_dict().get('appointment_date','') for s in fs_appointments().stream()}
        for d in docs: d['appointment_booked_on']=appts.get(d.get('po_number',''),'')
    return docs

def cloud_lines(doc_type):
    docs_by_id={d['id']:d for d in cloud_documents(doc_type)}
    out=[]
    for s in fs_lines().where('doc_type','==',doc_type.upper()).stream():
        x=s.to_dict(); d=docs_by_id.get(x.get('doc_id'))
        if not d or d.get('status')!='Success': continue
        row=dict(x.get('data') or {}); row['_line_id']=x.get('id',s.id); row['_doc_id']=x.get('doc_id'); row['_filename']=d.get('filename',''); row['_uploaded_at']=d.get('uploaded_at',''); out.append((x.get('row_index',0),row))
    out.sort(key=lambda z:(z[1].get('_uploaded_at',''),z[0]), reverse=False)
    return [r for _,r in out]

def local_documents(doc_type):
    c=conn(); rs=c.execute('SELECT * FROM documents WHERE doc_type=? ORDER BY uploaded_at DESC',(doc_type.upper(),)).fetchall(); appts={r['po_number']:r['appointment_date'] for r in c.execute('SELECT * FROM appointments').fetchall()}; c.close()
    out=[]
    for r in rs:
        d=dict(r)
        if doc_type.upper()=='PO': d['appointment_booked_on']=appts.get(d['po_number'],'')
        out.append(d)
    return out

def get_documents(doc_type): return cloud_documents(doc_type) if USE_FIRESTORE else local_documents(doc_type)

def get_lines(doc_type):
    if USE_FIRESTORE: return cloud_lines(doc_type)
    c=conn(); rs=c.execute('SELECT l.id AS line_id,l.data_json,l.doc_id,d.filename,d.uploaded_at FROM lines l JOIN documents d ON d.id=l.doc_id WHERE l.doc_type=? AND d.status="Success" ORDER BY d.uploaded_at,l.row_index',(doc_type.upper(),)).fetchall(); c.close()
    out=[]
    for r in rs:
        x=json.loads(r['data_json']); x['_line_id']=r['line_id']; x['_doc_id']=r['doc_id']; x['_filename']=r['filename']; x['_uploaded_at']=r['uploaded_at']; out.append(x)
    return out

def delete_line_data(line_id):
    if USE_FIRESTORE:
        ref=fs_lines().document(line_id); snap=ref.get()
        if not snap.exists: return None
        row=snap.to_dict(); ref.delete()
        remaining=list(fs_lines().where('doc_id','==',row['doc_id']).stream()); count=len(remaining)
        dref=fs_docs().document(row['doc_id']); d=dref.get()
        if d.exists: dref.update({'item_count':count, 'status':'Empty' if count==0 else 'Success','error':'All item rows deleted' if count==0 else ''})
        return {'id':line_id,'doc_id':row['doc_id'],'remaining_rows':count}
    c=conn(); row=c.execute('SELECT id,doc_id,doc_type FROM lines WHERE id=?',(line_id,)).fetchone()
    if not row: c.close(); return None
    c.execute('DELETE FROM lines WHERE id=?',(line_id,)); remaining=c.execute('SELECT COUNT(*) n FROM lines WHERE doc_id=?',(row['doc_id'],)).fetchone()['n']; c.execute('UPDATE documents SET item_count=?,status=?,error=? WHERE id=?',(remaining,'Empty' if remaining==0 else 'Success','All item rows deleted' if remaining==0 else '',row['doc_id'])); c.commit(); c.close(); return {'id':line_id,'doc_id':row['doc_id'],'remaining_rows':remaining}

def delete_document_data(did):
    if USE_FIRESTORE:
        dref=fs_docs().document(did); snap=dref.get()
        if not snap.exists: return None
        d=snap.to_dict(); lines=list(fs_lines().where('doc_id','==',did).stream())
        for s in lines: s.reference.delete()
        if d.get('doc_type')=='PO' and d.get('po_number'):
            others=[x for x in fs_docs().where('doc_type','==','PO').stream()
                    if x.id!=did and (x.to_dict() or {}).get('status')=='Success'
                    and (x.to_dict() or {}).get('po_number')==d['po_number']]
            if not others: fs_appointments().document(d['po_number']).delete()
        dref.delete()
        try:
            if d.get('storage_path'): _storage_bucket.blob(d['storage_path']).delete()
        except Exception: pass
        return {'id':did,'doc_type':d.get('doc_type'),'filename':d.get('filename')}
    c=conn(); d=c.execute('SELECT * FROM documents WHERE id=?',(did,)).fetchone()
    if not d: c.close(); return None
    c.execute('DELETE FROM lines WHERE doc_id=?',(did,))
    if d['doc_type']=='PO' and d['po_number']:
        other=c.execute('SELECT COUNT(*) n FROM documents WHERE doc_type="PO" AND status="Success" AND po_number=? AND id<>?',(d['po_number'],did)).fetchone()['n']
        if other==0: c.execute('DELETE FROM appointments WHERE po_number=?',(d['po_number'],))
    c.execute('DELETE FROM documents WHERE id=?',(did,)); c.commit(); c.close()
    try:
        if d['stored_name']: (ORIGINALS/d['stored_name']).unlink(missing_ok=True)
    except Exception: pass
    return {'id':did,'doc_type':d['doc_type'],'filename':d['filename']}

def set_appointment(po,date):
    now=datetime.now().isoformat(timespec='seconds')
    if USE_FIRESTORE:
        ref=fs_appointments().document(po)
        if date: ref.set({'po_number':po,'appointment_date':date,'updated_at':now})
        else: ref.delete()
    else:
        c=conn()
        if date: c.execute('INSERT INTO appointments(po_number,appointment_date,updated_at) VALUES(?,?,?) ON CONFLICT(po_number) DO UPDATE SET appointment_date=excluded.appointment_date,updated_at=excluded.updated_at',(po,date,now))
        else: c.execute('DELETE FROM appointments WHERE po_number=?',(po,))
        c.commit(); c.close()

def get_reconciliation():
    po_docs=[d for d in get_documents('PO') if d.get('status')=='Success']; grn_docs=[d for d in get_documents('GRN') if d.get('status')=='Success']
    po_lines={}; po_meta={}; grn_lines={}; grn_numbers={}
    def add_fin(bucket,key,row,qty_key,taxable_key,total_key):
        x=bucket.setdefault(key,{'qty':0.0,'taxable':0.0,'total':0.0}); x['qty']+=clean_num(row.get(qty_key)); x['taxable']+=clean_num(row.get(taxable_key)); x['total']+=clean_num(row.get(total_key))
    for d in po_docs:
        p=normalize_key(d.get('po_number')); po_meta[p]=dict(po_number=d.get('po_number'),vendor=d.get('vendor_name'),file_id=d.get('id'))
        for r in get_lines_for_doc(d['id'],'PO'): add_fin(po_lines,(p,normalize_key(r.get('Item Code'))),r,'Qty','Taxable Value (INR)','Total (INR)')
    for d in grn_docs:
        p=normalize_key(d.get('po_number')); grn_numbers.setdefault(p,[]).append(d.get('grn_number',''))
        for r in get_lines_for_doc(d['id'],'GRN'): add_fin(grn_lines,(p,normalize_key(r.get('SKU Code'))),r,'Recv Qty','Taxable Value (INR)','Total (INR)')
    if USE_FIRESTORE: appointments={s.id:s.to_dict().get('appointment_date','') for s in fs_appointments().stream()}
    else:
        c=conn(); appointments={r['po_number']:r['appointment_date'] for r in c.execute('SELECT * FROM appointments').fetchall()}; c.close()
    out=[]
    for key in sorted(set(po_lines)|set(grn_lines)):
        p,item=key; po=po_lines.get(key,{'qty':0.0,'taxable':0.0,'total':0.0}); gr=grn_lines.get(key,{'qty':0.0,'taxable':0.0,'total':0.0})
        ordered=po['qty']; received=gr['qty']; diff=ordered-received; po_tax=po['taxable']; grn_tax=gr['taxable']; tax_diff=po_tax-grn_tax; po_total=po['total']; grn_total=gr['total']; total_diff=po_total-grn_total
        if p not in po_meta: status='PO Not Found'
        elif key not in po_lines: status='Item Not in PO'
        elif received==0: status='Not Received'
        elif diff>0: status='Short Received'
        elif diff<0: status='Over Received'
        else: status='Matched'
        po_display=po_meta.get(p,{}).get('po_number',p)
        out.append({'po_number':po_display,'item_code':item,'ordered_qty':ordered,'received_qty':received,'difference_qty':diff,'po_taxable':po_tax,'grn_taxable':grn_tax,'taxable_difference':tax_diff,'po_total':po_total,'grn_total':grn_total,'total_difference':total_diff,'grn_numbers':', '.join(sorted(set(x for x in grn_numbers.get(p,[]) if x))),'appointment_booked_on':appointments.get(po_display,''),'status':status})
    return out

def get_lines_for_doc(did,doc_type):
    if USE_FIRESTORE:
        out=[]
        for s in fs_lines().where('doc_id','==',did).stream():
            x=s.to_dict()
            if x.get('doc_type')==doc_type: out.append((x.get('row_index',0), x.get('data') or {}))
        out.sort(key=lambda pair: pair[0])
        return [row for _, row in out]
    c=conn(); rs=c.execute('SELECT data_json FROM lines WHERE doc_id=? ORDER BY row_index',(did,)).fetchall(); c.close(); return [json.loads(x['data_json']) for x in rs]

# ---------- API ----------
@app.get('/')
def home(): return send_from_directory('static','index.html')
@app.get('/api/health')
def health(): return jsonify({'ok':True,'firestore':USE_FIRESTORE})

@app.post('/api/upload/<doc_type>')
@require_auth
def upload(doc_type):
    doc_type=doc_type.upper()
    if doc_type not in ('PO','GRN'): return jsonify({'error':'Invalid type'}),400
    files=request.files.getlist('files'); return jsonify([save_doc(doc_type,f) for f in files if f.filename])

@app.get('/api/documents/<doc_type>')
@require_auth
def documents(doc_type): return jsonify(get_documents(doc_type))

@app.get('/api/lines/<doc_type>')
@require_auth
def lines(doc_type): return jsonify(get_lines(doc_type))

@app.delete('/api/line/<line_id>')
@require_auth
def delete_line(line_id):
    result=delete_line_data(line_id)
    if not result: return jsonify({'error':'Row not found'}),404
    return jsonify({'ok':True,**result})

@app.delete('/api/document/<did>')
@require_auth
def delete_document(did):
    result=delete_document_data(did)
    if not result: return jsonify({'error':'Document not found'}),404
    return jsonify({'ok':True,**result})

@app.post('/api/appointment')
@require_auth
def appointment():
    data=request.get_json(force=True); po=str(data.get('po_number','')).strip(); date=str(data.get('appointment_booked_on','')).strip()
    if not po: return jsonify({'error':'PO Number is required'}),400
    set_appointment(po,date); return jsonify({'ok':True,'po_number':po,'appointment_booked_on':date})

@app.get('/api/reconciliation')
@require_auth
def reconciliation(): return jsonify(get_reconciliation())

@app.get('/api/stats')
@require_auth
def stats():
    po=[d for d in get_documents('PO') if d.get('status')=='Success']; gr=[d for d in get_documents('GRN') if d.get('status')=='Success']; rec=get_reconciliation()
    return jsonify({'po_files':len(po),'po_items':sum(int(d.get('item_count',0) or 0) for d in po),'grn_files':len(gr),'grn_items':sum(int(d.get('item_count',0) or 0) for d in gr),'matched':sum(x['status']=='Matched' for x in rec),'short':sum(x['status']=='Short Received' for x in rec),'over':sum(x['status']=='Over Received' for x in rec),'not_received':sum(x['status']=='Not Received' for x in rec),'po_not_found':sum(x['status']=='PO Not Found' for x in rec),'pending_qty':sum(max(0,x['difference_qty']) for x in rec)})

# ---------- exports ----------
def make_xlsx(headers,rows,name):
    wb=Workbook(); ws=wb.active; ws.title='Report'; ws.append(headers)
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78'); c.alignment=Alignment(wrap_text=True)
    for row in rows: ws.append([row.get(h,'') for h in headers])
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:
        width=min(45,max(10,max(len(str(c.value or '')) for c in col)+2)); ws.column_dimensions[col[0].column_letter].width=width
    p=GENERATED/name; wb.save(p); return p

def all_lines(t): return get_lines(t)
RECON_HEADERS=['PO Number','Item Code','Ordered Qty','Received Qty','Difference Qty','PO Taxable Value (INR)','GRN Taxable Value (INR)','Taxable Difference (INR)','PO Total (INR)','GRN Total (INR)','Total Difference (INR)','GRN Number(s)','Appointment Booked On','Status']
def recon_export_rows(rows):
    mapping={'PO Number':'po_number','Item Code':'item_code','Ordered Qty':'ordered_qty','Received Qty':'received_qty','Difference Qty':'difference_qty','PO Taxable Value (INR)':'po_taxable','GRN Taxable Value (INR)':'grn_taxable','Taxable Difference (INR)':'taxable_difference','PO Total (INR)':'po_total','GRN Total (INR)':'grn_total','Total Difference (INR)':'total_difference','GRN Number(s)':'grn_numbers','Appointment Booked On':'appointment_booked_on','Status':'status'}
    return [{h:r.get(k,'') for h,k in mapping.items()} for r in rows]

@app.get('/api/export/<kind>.xlsx')
@require_auth
def export_xlsx(kind):
    kind=kind.lower()
    if kind=='po': p=make_xlsx(PO_COLUMNS,all_lines('PO'),'PO_Consolidated.xlsx')
    elif kind=='grn': p=make_xlsx(GRN_COLUMNS,all_lines('GRN'),'GRN_Consolidated.xlsx')
    elif kind=='reconciliation': p=make_xlsx(RECON_HEADERS,recon_export_rows(get_reconciliation()),'PO_GRN_Reconciliation.xlsx')
    else:return 'Invalid',400
    return send_file(p,as_attachment=True,download_name=p.name)

@app.post('/api/export/reconciliation-visible.xlsx')
@require_auth
def export_reconciliation_visible_xlsx():
    payload=request.get_json(silent=True) or {}; p=make_xlsx(RECON_HEADERS,recon_export_rows(payload.get('rows') or []),'PO_GRN_Reconciliation_Visible.xlsx'); return send_file(p,as_attachment=True,download_name=p.name)

@app.get('/api/document/<did>/original')
@require_auth
def original(did):
    if USE_FIRESTORE:
        dref=fs_docs().document(did).get()
        if not dref.exists or not dref.to_dict().get('storage_path'): return 'Not found',404
        d=dref.to_dict(); data=_storage_bucket.blob(d['storage_path']).download_as_bytes(); return send_file(io.BytesIO(data),as_attachment=True,download_name=d.get('filename','document.pdf'))
    c=conn(); r=c.execute('SELECT * FROM documents WHERE id=?',(did,)).fetchone(); c.close()
    if not r or not r['stored_name']: return 'Not found',404
    return send_file(ORIGINALS/r['stored_name'],as_attachment=True,download_name=r['filename'])

@app.get('/api/document/<did>/excel')
@require_auth
def doc_excel(did):
    docs=get_documents('PO')+get_documents('GRN'); d=next((x for x in docs if x.get('id')==did),None)
    if not d:return 'Not found',404
    rows=get_lines_for_doc(did,d['doc_type']); headers=PO_COLUMNS if d['doc_type']=='PO' else GRN_COLUMNS; name=f"{d.get('po_number') or d.get('grn_number') or did}_{d['doc_type']}.xlsx"; p=make_xlsx(headers,rows,name); return send_file(p,as_attachment=True,download_name=p.name)

def pdf_table(title,headers,rows,path):
    styles=getSampleStyleSheet(); doc=SimpleDocTemplate(str(path),pagesize=landscape(A4),leftMargin=16,rightMargin=16,topMargin=16,bottomMargin=16); data=[[str(x) for x in headers]]+[[str(r.get(x,'')) for x in headers] for r in rows]; t=Table(data,repeatRows=1); t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#1F4E78')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.25,colors.grey),('FONTSIZE',(0,0),(-1,-1),5.5),('VALIGN',(0,0),(-1,-1),'TOP')])); doc.build([Paragraph(title,styles['Title']),Spacer(1,8),t]); return path

@app.get('/api/export/po.pdf')
@require_auth
def po_pdf(): p=GENERATED/'PO_Consolidated.pdf'; pdf_table('Purchase Orders - Consolidated',PO_COLUMNS,all_lines('PO'),p); return send_file(p,as_attachment=True,download_name=p.name)
@app.get('/api/export/grn.pdf')
@require_auth
def grn_pdf(): p=GENERATED/'GRN_Consolidated.pdf'; pdf_table('GRNs - Consolidated',GRN_COLUMNS,all_lines('GRN'),p); return send_file(p,as_attachment=True,download_name=p.name)
@app.get('/api/export/reconciliation.pdf')
@require_auth
def recon_pdf(): p=GENERATED/'PO_GRN_Reconciliation.pdf'; pdf_table('PO ↔ GRN Reconciliation',RECON_HEADERS,recon_export_rows(get_reconciliation()),p); return send_file(p,as_attachment=True,download_name=p.name)
@app.post('/api/export/reconciliation-visible.pdf')
@require_auth
def export_reconciliation_visible_pdf(): payload=request.get_json(silent=True) or {}; p=GENERATED/'PO_GRN_Reconciliation_Visible.pdf'; pdf_table('PO ↔ GRN Reconciliation - Visible Rows',RECON_HEADERS,recon_export_rows(payload.get('rows') or []),p); return send_file(p,as_attachment=True,download_name=p.name)
@app.get('/api/document/<did>/pdf')
@require_auth
def doc_pdf(did):
    docs=get_documents('PO')+get_documents('GRN'); d=next((x for x in docs if x.get('id')==did),None)
    if not d:return 'Not found',404
    rows=get_lines_for_doc(did,d['doc_type']); headers=PO_COLUMNS if d['doc_type']=='PO' else GRN_COLUMNS; base=d.get('po_number') or d.get('grn_number') or did; p=GENERATED/f'{base}_{d["doc_type"]}.pdf'; pdf_table(f'{d["doc_type"]} {base}',headers,rows,p); return send_file(p,as_attachment=True,download_name=p.name)

if __name__=='__main__': app.run(host='127.0.0.1',port=3002,debug=False)
